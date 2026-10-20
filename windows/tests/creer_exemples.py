"""Cree les documents d'exemple des tests de reperes (aucune donnee personnelle)."""
import zipfile, io
from fpdf import FPDF
from pypdf import PdfReader, PdfWriter
import docx
from docx.enum.text import WD_BREAK

CORPS = "Le village se reveilla sous la neige. Marc ouvrit la porte et regarda le lac gele. "
# PDF : 2 pages liminaires (i, ii) puis 4 pages (1 a 4), signets sur deux niveaux
pdf = FPDF(); pdf.set_font('Helvetica', size=12)
pdf.add_page(); pdf.multi_cell(0, 8, "Avant-propos\n" + CORPS * 3)
pdf.add_page(); pdf.multi_cell(0, 8, "Table des matieres\nChapitre 1 La neige ........ 1\nChapitre 2 Le lac ........ 3")
pdf.add_page(); pdf.start_section("Chapitre 1 La neige", level=0); pdf.multi_cell(0, 8, "Chapitre 1 La neige\n" + CORPS * 5)
pdf.add_page(); pdf.start_section("La cabane", level=1); pdf.multi_cell(0, 8, CORPS + "\nLa cabane\n" + CORPS * 3)
pdf.add_page(); pdf.start_section("Chapitre 2 Le lac", level=0); pdf.multi_cell(0, 8, "Chapitre 2 Le lac\n" + CORPS * 4)
pdf.add_page(); pdf.multi_cell(0, 8, "Fin du livre. " + CORPS)
brut = io.BytesIO(pdf.output())
w = PdfWriter(clone_from=PdfReader(brut))
w.set_page_label(0, 1, style='/r')
w.set_page_label(2, 5, style='/D', start=1)
w.write('livre.pdf')
# meme livre sans signets (chapitres devines)
pdf2 = FPDF(); pdf2.set_font('Helvetica', size=12)
for titre in ("Introduction", "Chapitre 1 La neige", "Chapitre 2 Le lac", "Conclusion"):
    pdf2.add_page(); pdf2.multi_cell(0, 8, titre + "\n" + CORPS * 3)
pdf2.output('sans-signets.pdf')
# DOCX : titres + sauts de page
d = docx.Document()
d.add_heading('Premiere partie', level=1); d.add_paragraph(CORPS * 2)
d.add_heading('Le depart', level=2); d.add_paragraph(CORPS)
d.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
d.add_paragraph(CORPS); d.add_heading('Deuxieme partie', level=1); d.add_paragraph(CORPS)
d.save('document.docx')
# ODT minimal : titres + saut de page doux
contenu = '''<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2"><office:body><office:text>
<text:h text:outline-level="1">Premier titre</text:h><text:p>''' + CORPS + '''</text:p>
<text:soft-page-break/><text:h text:outline-level="2">Sous-titre</text:h><text:p>''' + CORPS + '''</text:p>
<text:h text:outline-level="1">Second titre</text:h><text:p>Fin.</text:p></office:text></office:body></office:document-content>'''
with zipfile.ZipFile('document.odt', 'w') as z:
    z.writestr('mimetype', 'application/vnd.oasis.opendocument.text'); z.writestr('content.xml', contenu)
# EPUB 3 (nav) et EPUB 2 (ncx)
def epub(nom, nav):
    with zipfile.ZipFile(nom, 'w') as z:
        z.writestr('mimetype', 'application/epub+zip')
        z.writestr('META-INF/container.xml', '<?xml version="1.0"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0"><rootfiles><rootfile full-path="OEBPS/livre.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
        man = '<item id="c1" href="texte/ch1.xhtml" media-type="application/xhtml+xml"/><item id="c2" href="texte/ch2.xhtml" media-type="application/xhtml+xml"/>'
        man += '<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>' if nav else '<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>'
        z.writestr('OEBPS/livre.opf', '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0"><manifest>' + man + '</manifest><spine><itemref idref="c1"/><itemref idref="c2"/></spine></package>')
        for i, t in ((1, 'Le départ'), (2, 'Le retour')):
            z.writestr('OEBPS/texte/ch%d.xhtml' % i, '<html xmlns="http://www.w3.org/1999/xhtml"><body><h1>Chapitre %d : %s</h1><p>%s</p><h2>Le soir</h2><p>%s</p></body></html>' % (i, t, CORPS * 2, CORPS))
        if nav:
            z.writestr('OEBPS/nav.xhtml', '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><body><nav epub:type="toc"><ol><li><a href="texte/ch1.xhtml">Chapitre 1 : Le départ</a><ol><li><a href="texte/ch1.xhtml#s">Le soir</a></li></ol></li><li><a href="texte/ch2.xhtml">Chapitre 2 : Le retour</a></li></ol></nav></body></html>')
        else:
            z.writestr('OEBPS/toc.ncx', '<?xml version="1.0"?><ncx xmlns="http://www.daisy.org/z3986/2005/ncx/"><navMap><navPoint id="a"><navLabel><text>Chapitre 1 : Le départ</text></navLabel><content src="texte/ch1.xhtml"/></navPoint><navPoint id="b"><navLabel><text>Chapitre 2 : Le retour</text></navLabel><content src="texte/ch2.xhtml"/></navPoint></navMap></ncx>')
epub('livre3.epub', True); epub('livre2.epub', False)
open('notes.md', 'w', encoding='utf-8').write('# Mon carnet\n\nIntro.\n\n## Lundi\n\nIl a neigé.\n\n## Mardi\n\nSoleil. \U0001F600\n')
print('exemples crees')
